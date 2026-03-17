import io
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


class ExportService:

    @staticmethod
    def _style_header_row(ws, col_count: int) -> None:
        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_fill = PatternFill(start_color="2B579A", end_color="2B579A", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )
        for col in range(1, col_count + 1):
            cell = ws.cell(row=1, column=col)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border

    @staticmethod
    def export_bookings_excel(
        bookings: list[dict],
        date_from: date | None = None,
        date_to: date | None = None,
    ) -> io.BytesIO:
        wb = Workbook()
        ws = wb.active
        ws.title = "Бронирования"

        headers = [
            "№", "Дата", "Время", "Направление", "Ресурс",
            "Студент", "Email", "Статус", "Посещение", "Создано",
        ]
        ws.append(headers)
        ExportService._style_header_row(ws, len(headers))

        for i, b in enumerate(bookings, 1):
            ws.append([
                i,
                b.get("date", ""),
                b.get("time", ""),
                b.get("direction", ""),
                b.get("resource", ""),
                b.get("student", ""),
                b.get("email", ""),
                b.get("status", ""),
                b.get("attended", ""),
                b.get("created_at", ""),
            ])

        for col in ws.columns:
            max_length = 0
            column_letter = col[0].column_letter
            for cell in col:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            ws.column_dimensions[column_letter].width = min(max_length + 4, 40)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    @staticmethod
    def export_users_excel(users: list[dict]) -> io.BytesIO:
        wb = Workbook()
        ws = wb.active
        ws.title = "Пользователи"

        headers = [
            "№", "Фамилия", "Имя", "Отчество", "Email",
            "Роль", "Рейтинг", "Верифицирован", "Активен", "Дата регистрации",
        ]
        ws.append(headers)
        ExportService._style_header_row(ws, len(headers))

        for i, u in enumerate(users, 1):
            ws.append([
                i,
                u.get("last_name", ""),
                u.get("first_name", ""),
                u.get("patronymic", ""),
                u.get("email", ""),
                u.get("role", ""),
                u.get("rating_score", 0),
                "Да" if u.get("is_verified") else "Нет",
                "Да" if u.get("is_active") else "Нет",
                u.get("created_at", ""),
            ])

        for col in ws.columns:
            max_length = 0
            column_letter = col[0].column_letter
            for cell in col:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            ws.column_dimensions[column_letter].width = min(max_length + 4, 40)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    @staticmethod
    def export_rating_excel(rating_data: list[dict]) -> io.BytesIO:
        wb = Workbook()
        ws = wb.active
        ws.title = "Рейтинг"

        headers = [
            "Место", "Фамилия", "Имя", "Email", "Рейтинг",
        ]
        ws.append(headers)
        ExportService._style_header_row(ws, len(headers))

        for i, r in enumerate(rating_data, 1):
            ws.append([
                i,
                r.get("last_name", ""),
                r.get("first_name", ""),
                r.get("email", ""),
                r.get("rating_score", 0),
            ])

        for col in ws.columns:
            max_length = 0
            column_letter = col[0].column_letter
            for cell in col:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            ws.column_dimensions[column_letter].width = min(max_length + 4, 40)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    @staticmethod
    def export_bookings_pdf(
        bookings: list[dict],
        date_from: date | None = None,
        date_to: date | None = None,
    ) -> io.BytesIO:
        output = io.BytesIO()
        doc = SimpleDocTemplate(
            output,
            pagesize=landscape(A4),
            leftMargin=15 * mm,
            rightMargin=15 * mm,
            topMargin=15 * mm,
            bottomMargin=15 * mm,
        )

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            "CustomTitle",
            parent=styles["Title"],
            fontSize=16,
            spaceAfter=12,
        )

        elements = []

        title_text = "Отчёт по бронированиям"
        if date_from and date_to:
            title_text += f" ({date_from.isoformat()} — {date_to.isoformat()})"
        elif date_from:
            title_text += f" (с {date_from.isoformat()})"
        elif date_to:
            title_text += f" (до {date_to.isoformat()})"

        elements.append(Paragraph(title_text, title_style))
        elements.append(Spacer(1, 6 * mm))

        table_data = [
            ["№", "Дата", "Время", "Направление", "Студент", "Статус", "Посещение"],
        ]

        cell_style = ParagraphStyle("Cell", fontSize=8, leading=10)

        for i, b in enumerate(bookings, 1):
            table_data.append([
                str(i),
                str(b.get("date", "")),
                str(b.get("time", "")),
                Paragraph(str(b.get("direction", "")), cell_style),
                Paragraph(str(b.get("student", "")), cell_style),
                str(b.get("status", "")),
                str(b.get("attended", "")),
            ])

        col_widths = [25, 60, 60, 120, 150, 70, 60]

        table = Table(table_data, colWidths=col_widths, repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2B579A")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, 0), 9),
            ("FONTSIZE", (0, 1), (-1, -1), 8),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F2F2")]),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))

        elements.append(table)
        elements.append(Spacer(1, 6 * mm))

        generated_text = f"Сформировано: {datetime.now().strftime('%d.%m.%Y %H:%M')}"
        elements.append(Paragraph(generated_text, styles["Normal"]))

        doc.build(elements)
        output.seek(0)
        return output
